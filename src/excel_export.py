from datetime import datetime
from pathlib import Path
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .database import get_receipts_by_month, get_receipt_items, get_monthly_summary

EXPORTS_PATH = Path(__file__).parent.parent / "exports"


def style_header(cell):
    """Stijl een header cel."""
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center")


def style_currency(cell):
    """Stijl een valuta cel."""
    cell.number_format = '€ #,##0.00'
    cell.alignment = Alignment(horizontal="right")


def add_border(cell):
    """Voeg een rand toe aan een cel."""
    thin = Side(style='thin', color="CCCCCC")
    cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)


async def create_monthly_report(
    telegram_user_id: int,
    year: int,
    month: int
) -> Optional[str]:
    """
    Maak een Excel rapport voor een specifieke maand.
    Retourneert het pad naar het bestand.
    """
    # Haal data op
    receipts = await get_receipts_by_month(telegram_user_id, year, month)
    summary = await get_monthly_summary(telegram_user_id, year, month)

    if not receipts:
        return None

    # Maak exports directory als die niet bestaat
    EXPORTS_PATH.mkdir(exist_ok=True)

    # Maak workbook
    wb = Workbook()

    # === OVERZICHT SHEET ===
    ws_overview = wb.active
    ws_overview.title = "Overzicht"

    # Titel
    ws_overview.merge_cells('A1:D1')
    title_cell = ws_overview['A1']
    title_cell.value = f"Bonnetjes Overzicht - {month:02d}/{year}"
    title_cell.font = Font(bold=True, size=16)
    title_cell.alignment = Alignment(horizontal="center")

    # Samenvatting sectie
    ws_overview['A3'] = "Samenvatting"
    ws_overview['A3'].font = Font(bold=True, size=12)

    summary_data = [
        ("Totaal aantal bonnetjes:", summary.get('total_receipts', 0)),
        ("Totaal uitgegeven:", summary.get('total_spent', 0)),
        ("Totaal BTW:", summary.get('total_btw', 0)),
        ("Totaal PIN:", summary.get('total_pin', 0)),
        ("Totaal Cash:", summary.get('total_cash', 0)),
        ("Totaal Onbekend:", summary.get('total_unknown', 0)),
    ]

    for i, (label, value) in enumerate(summary_data, start=4):
        ws_overview[f'A{i}'] = label
        ws_overview[f'B{i}'] = value if isinstance(value, int) and i == 4 else (value or 0)
        if i > 4:  # Valuta cellen
            style_currency(ws_overview[f'B{i}'])

    # === BONNETJES SHEET ===
    ws_receipts = wb.create_sheet("Bonnetjes")

    # Headers
    headers = ["ID", "Datum", "Winkel", "Categorie", "Totaal", "BTW", "Betaalmethode"]
    for col, header in enumerate(headers, start=1):
        cell = ws_receipts.cell(row=1, column=col, value=header)
        style_header(cell)
        add_border(cell)

    # Data
    for row, receipt in enumerate(receipts, start=2):
        ws_receipts.cell(row=row, column=1, value=receipt['id'])
        ws_receipts.cell(row=row, column=2, value=receipt.get('date', ''))
        ws_receipts.cell(row=row, column=3, value=receipt.get('store_name', ''))
        ws_receipts.cell(row=row, column=4, value=receipt.get('category', ''))

        total_cell = ws_receipts.cell(row=row, column=5, value=receipt.get('total_amount', 0))
        style_currency(total_cell)

        btw_cell = ws_receipts.cell(row=row, column=6, value=receipt.get('btw_amount', 0))
        style_currency(btw_cell)

        ws_receipts.cell(row=row, column=7, value=receipt.get('payment_method', '').upper())

        # Borders toevoegen
        for col in range(1, 8):
            add_border(ws_receipts.cell(row=row, column=col))

    # Kolom breedtes aanpassen
    ws_receipts.column_dimensions['A'].width = 8
    ws_receipts.column_dimensions['B'].width = 12
    ws_receipts.column_dimensions['C'].width = 25
    ws_receipts.column_dimensions['D'].width = 15
    ws_receipts.column_dimensions['E'].width = 12
    ws_receipts.column_dimensions['F'].width = 10
    ws_receipts.column_dimensions['G'].width = 15

    # Totaal rij
    total_row = len(receipts) + 2
    ws_receipts.cell(row=total_row, column=4, value="TOTAAL:").font = Font(bold=True)

    total_sum_cell = ws_receipts.cell(
        row=total_row, column=5,
        value=f"=SUM(E2:E{total_row-1})"
    )
    style_currency(total_sum_cell)
    total_sum_cell.font = Font(bold=True)

    btw_sum_cell = ws_receipts.cell(
        row=total_row, column=6,
        value=f"=SUM(F2:F{total_row-1})"
    )
    style_currency(btw_sum_cell)
    btw_sum_cell.font = Font(bold=True)

    # === PER CATEGORIE SHEET ===
    ws_category = wb.create_sheet("Per Categorie")

    # Headers
    cat_headers = ["Categorie", "Aantal", "Totaal", "Gemiddeld"]
    for col, header in enumerate(cat_headers, start=1):
        cell = ws_category.cell(row=1, column=col, value=header)
        style_header(cell)

    # Bereken per categorie
    categories = {}
    for receipt in receipts:
        cat = receipt.get('category', 'overig') or 'overig'
        if cat not in categories:
            categories[cat] = {'count': 0, 'total': 0}
        categories[cat]['count'] += 1
        categories[cat]['total'] += receipt.get('total_amount', 0) or 0

    for row, (cat, data) in enumerate(sorted(categories.items()), start=2):
        ws_category.cell(row=row, column=1, value=cat.capitalize())
        ws_category.cell(row=row, column=2, value=data['count'])

        total_cell = ws_category.cell(row=row, column=3, value=data['total'])
        style_currency(total_cell)

        avg_cell = ws_category.cell(
            row=row, column=4,
            value=data['total'] / data['count'] if data['count'] > 0 else 0
        )
        style_currency(avg_cell)

    # Kolom breedtes
    ws_category.column_dimensions['A'].width = 15
    ws_category.column_dimensions['B'].width = 10
    ws_category.column_dimensions['C'].width = 12
    ws_category.column_dimensions['D'].width = 12

    # === PER BETAALMETHODE SHEET ===
    ws_payment = wb.create_sheet("Per Betaalmethode")

    # Headers
    payment_headers = ["Betaalmethode", "Aantal", "Totaal"]
    for col, header in enumerate(payment_headers, start=1):
        cell = ws_payment.cell(row=1, column=col, value=header)
        style_header(cell)

    # Bereken per betaalmethode
    payment_methods = {}
    for receipt in receipts:
        method = receipt.get('payment_method', 'onbekend') or 'onbekend'
        if method not in payment_methods:
            payment_methods[method] = {'count': 0, 'total': 0}
        payment_methods[method]['count'] += 1
        payment_methods[method]['total'] += receipt.get('total_amount', 0) or 0

    for row, (method, data) in enumerate(sorted(payment_methods.items()), start=2):
        ws_payment.cell(row=row, column=1, value=method.upper())
        ws_payment.cell(row=row, column=2, value=data['count'])

        total_cell = ws_payment.cell(row=row, column=3, value=data['total'])
        style_currency(total_cell)

    # Kolom breedtes
    ws_payment.column_dimensions['A'].width = 15
    ws_payment.column_dimensions['B'].width = 10
    ws_payment.column_dimensions['C'].width = 12

    # Sla bestand op
    filename = f"bonnetjes_{year}_{month:02d}_{telegram_user_id}.xlsx"
    filepath = EXPORTS_PATH / filename
    wb.save(filepath)

    return str(filepath)
